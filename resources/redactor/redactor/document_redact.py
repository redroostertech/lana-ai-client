"""TRUE document redaction: ingest PDF/DOCX, redact PII, emit a redacted file.

This is the "redact this document" product surface. It does NOT reimplement any
detection logic — it CALLS ``engine.redact`` (via an injected ``redact_fn``) on
the text pulled out of each document, then maps the returned entity CHAR SPANS
back onto document geometry and destroys the underlying sensitive content.

Two output guarantees, by format:

* **PDF** — real burn-in. We locate each redacted entity's char span in the
  page text (reconstructed from ``page.get_text("words")`` so char offsets and
  word rectangles share one coordinate space), collect the covering word
  rectangles, and apply ``page.add_redact_annot(rect)`` +
  ``page.apply_redactions()``. PyMuPDF's ``apply_redactions`` DELETES the glyphs
  (and image pixels) under each rectangle, so the sensitive strings are truly
  UN-EXTRACTABLE afterward — not merely covered by an opaque box.
* **DOCX** — there is no glyph-level burn-in in the OOXML text model, so we
  REPLACE the run text with the tokenized/blacked output of ``redact``. The
  original characters are overwritten in the part XML, not overlaid. We rewrite
  EVERY text-bearing region we have a safe API for — body paragraphs + tables
  (recursively), every section header/footer variant, and core-property
  metadata — and FAIL CLOSED on regions we cannot safely rewrite (comments,
  footnotes, endnotes, textboxes): if such a region carries real text we raise
  and emit NOTHING rather than an incompletely-redacted document.

Scanned / image-only PDF pages (no extractable text) are OCR'd with pytesseract
when the ``tesseract`` binary is available; the OCR'd words feed the same span ->
rectangle -> burn-in path. When OCR is unavailable we say so explicitly in the
report rather than emitting a silently-unredacted page.

Fail-closed: if ``redact_fn`` raises ``RedactionUnavailable`` (strict mode, NER
layer down) we let it propagate and emit NOTHING. A partially-redacted document
is worse than no document.
"""

from __future__ import annotations

import io
import logging
from collections import defaultdict
from dataclasses import dataclass
from typing import Any, Callable, Protocol

_log = logging.getLogger("redactor")


class _RedactResult(Protocol):
    redacted: str
    token_map: dict[str, str]
    entities: list[dict[str, Any]]


# A redact_fn takes page/paragraph text and returns a RedactionResult-shaped
# object. The service binds engine.redact with its loaded analyzer/PHI models;
# tests can pass engine.redact directly (with or without an analyzer).
RedactFn = Callable[[str], _RedactResult]

# A DetectFn takes text and returns the DETECT-ONLY resolved span list (no token
# emission) — the shape engine.detect produces: each span carries at least
# {type, start, end, score, source, text}. This is the PROPOSE side of the
# human-in-the-loop flow; nothing is burned in until apply_document runs on the
# human-approved subset.
DetectFn = Callable[[str], list[dict[str, Any]]]


class UnsupportedDocument(ValueError):
    """The uploaded bytes are not a format we can redact (not PDF/DOCX)."""


class PartialRedactionError(RuntimeError):
    """A DOCX carries text in a region we have no safe API to rewrite
    (comments / footnotes / endnotes / textboxes). We FAIL CLOSED and emit
    NOTHING rather than a partially-redacted "safe for lawyers" document."""


@dataclass
class DocumentRedactionResult:
    """Redacted document bytes + a per-entity redaction report."""

    content: bytes
    fmt: str  # "pdf" | "docx"
    report: dict[str, Any]


# --- OCR availability -------------------------------------------------------
def ocr_available() -> bool:
    """True only if BOTH the pytesseract wrapper and the tesseract binary exist.

    pytesseract is a thin wrapper around the ``tesseract`` CLI; importing it
    succeeds even when the binary is absent, so we probe the binary explicitly.
    """
    try:
        import pytesseract  # noqa: F401
    except Exception:
        return False
    try:
        pytesseract.get_tesseract_version()
        return True
    except Exception:
        return False


# --- report helpers ---------------------------------------------------------
def _new_report(fmt: str) -> dict[str, Any]:
    return {
        "format": fmt,
        "entity_counts": {},
        "total_entities": 0,
        "locations": [],  # per unit: {page|paragraph|cell, type, ...}
        "ocr": {"used": False, "available": ocr_available(), "pages": []},
    }


def _tally(report: dict[str, Any], etype: str) -> None:
    report["entity_counts"][etype] = report["entity_counts"].get(etype, 0) + 1
    report["total_entities"] += 1


# --- PDF --------------------------------------------------------------------
def _build_page_text_from_words(
    words: list[tuple],
) -> tuple[str, list[tuple[int, int, Any]]]:
    """Reconstruct a page's text from ``get_text('words')`` output.

    Returns ``(page_text, spans)`` where ``spans`` is a list of
    ``(char_start, char_end, rect)`` — the char range each word occupies in
    ``page_text`` and its rectangle. Because we build ``page_text`` ourselves,
    the char offsets that ``redact`` returns line up EXACTLY with these word
    spans; no fuzzy string search is needed.

    Word tuples are ``(x0, y0, x1, y1, text, block_no, line_no, word_no)`` in
    reading order. We join words on the same line with a single space and start
    a newline when the (block, line) changes, so the reconstructed text reads
    naturally for the detector (labels stay next to their values).
    """
    import fitz

    parts: list[str] = []
    spans: list[tuple[int, int, Any]] = []
    cursor = 0
    prev_key: tuple[int, int] | None = None
    for w in words:
        x0, y0, x1, y1, wtext = w[0], w[1], w[2], w[3], w[4]
        block_no, line_no = w[5], w[6]
        key = (block_no, line_no)
        if prev_key is not None:
            sep = " " if key == prev_key else "\n"
            parts.append(sep)
            cursor += len(sep)
        start = cursor
        parts.append(wtext)
        cursor += len(wtext)
        spans.append((start, cursor, fitz.Rect(x0, y0, x1, y1)))
        prev_key = key
    return "".join(parts), spans


def _rects_for_entity(
    ent: dict[str, Any], word_spans: list[tuple[int, int, Any]]
) -> list[Any]:
    """Word rectangles overlapping an entity's char span [start, end).

    A word overlaps the entity if ``ws < ee and es < we``. We operate at WORD
    granularity: if a redacted span covers only part of a word, the WHOLE word
    is burned. This errs toward OVER-redaction (a label char or trailing comma
    may be blacked too) but never under-redacts — the safe direction for a legal
    product.
    """
    es, ee = ent["start"], ent["end"]
    return [rect for ws, we, rect in word_spans if ws < ee and es < we]


def _ocr_words(page, zoom: float = 3.0) -> list[tuple]:
    """OCR a page into ``get_text('words')``-shaped tuples via pytesseract.

    Renders the page at ``zoom`` (higher = better OCR, at more pixels), runs
    ``image_to_data``, and maps pixel boxes back into PDF point space by dividing
    by ``zoom``. block/line come from tesseract's own block/line numbering so the
    line-join heuristic in ``_build_page_text_from_words`` still works.
    """
    import fitz
    import pytesseract
    from PIL import Image

    mat = fitz.Matrix(zoom, zoom)
    pix = page.get_pixmap(matrix=mat, alpha=False)
    img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
    data = pytesseract.image_to_data(img, output_type=pytesseract.Output.DICT)
    words: list[tuple] = []
    n = len(data["text"])
    for i in range(n):
        txt = (data["text"][i] or "").strip()
        if not txt:
            continue
        try:
            conf = float(data["conf"][i])
        except (TypeError, ValueError):
            conf = -1.0
        if conf < 0:  # tesseract emits -1 for non-text layout boxes
            continue
        left, top = data["left"][i], data["top"][i]
        w, h = data["width"][i], data["height"][i]
        x0, y0 = left / zoom, top / zoom
        x1, y1 = (left + w) / zoom, (top + h) / zoom
        block_no = data["block_num"][i]
        line_no = data["line_num"][i]
        word_no = data["word_num"][i]
        words.append((x0, y0, x1, y1, txt, block_no, line_no, word_no))
    return words


def redact_pdf(data: bytes, redact_fn: RedactFn) -> DocumentRedactionResult:
    """Burn-in redaction of a PDF. Sensitive text is DESTROYED, not covered."""
    import fitz

    report = _new_report("pdf")
    doc = fitz.open(stream=data, filetype="pdf")
    try:
        for pno in range(doc.page_count):
            page = doc.load_page(pno)
            words = page.get_text("words")
            used_ocr = False
            if not words:
                # No extractable text: image-only / scanned page.
                if ocr_available():
                    words = _ocr_words(page)
                    used_ocr = True
                    report["ocr"]["used"] = True
                    report["ocr"]["pages"].append({"page": pno, "ocr": True})
                else:
                    report["ocr"]["pages"].append(
                        {"page": pno, "ocr": False, "reason": "ocr_unavailable"}
                    )
                    _log.warning(
                        "redactor.doc.pdf_page_no_text_ocr_unavailable page=%d", pno
                    )
                    continue
            if not words:
                continue

            page_text, word_spans = _build_page_text_from_words(words)
            result = redact_fn(page_text)  # may raise RedactionUnavailable
            if not result.entities:
                continue

            for ent in result.entities:
                rects = _rects_for_entity(ent, word_spans)
                if not rects:
                    continue
                _tally(report, ent["type"])
                loc_rects = []
                for rect in rects:
                    page.add_redact_annot(rect, fill=(0, 0, 0))
                    loc_rects.append([rect.x0, rect.y0, rect.x1, rect.y1])
                report["locations"].append(
                    {"page": pno, "type": ent["type"], "ocr": used_ocr,
                     "rects": loc_rects}
                )
            # DESTROY the underlying content under every annot on this page.
            page.apply_redactions()

        out = io.BytesIO()
        doc.save(out, garbage=4, deflate=True, clean=True)
        return DocumentRedactionResult(out.getvalue(), "pdf", report)
    finally:
        doc.close()


# --- DOCX -------------------------------------------------------------------
def _redact_paragraph(paragraph, redact_fn: RedactFn, report: dict[str, Any],
                      location: dict[str, Any]) -> None:
    """Replace a paragraph's text with its redacted (tokenized) form.

    We redact the FULL paragraph text (so an entity spanning multiple runs is
    caught), then write the result into the first run and blank the rest. The
    first run's formatting is preserved; the original characters are OVERWRITTEN
    in the part XML, not overlaid. Empty / whitespace-only paragraphs are left
    untouched.
    """
    text = paragraph.text
    if not text.strip():
        return
    result = redact_fn(text)  # may raise RedactionUnavailable
    if not result.entities:
        return
    for ent in result.entities:
        _tally(report, ent["type"])
        report["locations"].append({**location, "type": ent["type"]})
    runs = paragraph.runs
    if runs:
        runs[0].text = result.redacted
        for r in runs[1:]:
            r.text = ""
    else:
        # No runs (rare): set via the paragraph-level add_run.
        paragraph.add_run(result.redacted)


# Core-property fields that can carry free-text PII (case caption, client name,
# matter number). datetime-valued props (created/modified/...) are skipped.
_METADATA_FIELDS: tuple[str, ...] = (
    "author", "last_modified_by", "title", "subject",
    "keywords", "comments", "category",
)

# DOCX parts that can carry free text for which python-docx exposes no
# first-class rewrite API. Presence WITH real text triggers fail-closed.
_UNHANDLED_TEXT_PARTS: tuple[str, ...] = (
    "comments.xml", "footnotes.xml", "endnotes.xml",
)

_W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_W_T = _W_NS + "t"
_W_TXBX = _W_NS + "txbxContent"


def _subtree_has_text(element) -> bool:
    """True if any ``w:t`` descendant carries non-whitespace text.

    Default footnote/endnote separators contain runs but no ``w:t`` text, so an
    auto-generated (empty) footnotes.xml correctly reads as text-free here."""
    for t in element.iter(_W_T):
        if (t.text or "").strip():
            return True
    return False


def _docx_unhandled_text_regions(data: bytes) -> list[str]:
    """Names of text-bearing DOCX regions we cannot safely rewrite.

    Scans the raw OOXML zip (independent of python-docx part registration) for
    comments/footnotes/endnotes parts and for textbox content
    (``w:txbxContent``, which python-docx does not surface through
    ``document.paragraphs``) that hold any non-empty ``w:t``."""
    import zipfile

    from lxml import etree

    found: list[str] = []
    try:
        zf = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile:
        return found
    with zf:
        for name in zf.namelist():
            if not name.endswith(".xml"):
                continue
            base = name.rsplit("/", 1)[-1]
            try:
                root = etree.fromstring(zf.read(name))
            except etree.XMLSyntaxError:
                continue
            if base in _UNHANDLED_TEXT_PARTS and _subtree_has_text(root):
                found.append(base)
                continue
            for tb in root.iter(_W_TXBX):
                if _subtree_has_text(tb):
                    found.append(f"textbox:{name}")
                    break
    return found


def _redact_metadata(document, redact_fn: RedactFn, report: dict[str, Any],
                     regions: set[str]) -> None:
    """Redact free-text core properties in place (replace with redacted text)."""
    cp = document.core_properties
    for field in _METADATA_FIELDS:
        val = getattr(cp, field, None)
        if not isinstance(val, str) or not val.strip():
            continue
        result = redact_fn(val)  # may raise RedactionUnavailable
        if not result.entities:
            continue
        for ent in result.entities:
            _tally(report, ent["type"])
            report["locations"].append(
                {"region": "metadata", "field": field, "type": ent["type"]}
            )
        setattr(cp, field, result.redacted)
        regions.add("metadata")


def redact_docx(data: bytes, redact_fn: RedactFn) -> DocumentRedactionResult:
    """Redact a DOCX by REPLACING run text (no overlay), across ALL text-bearing
    regions we can safely rewrite:

      * body paragraphs and table cells (recursively -> nested tables),
      * every section's header/footer variant (default, first-page, even-page),
      * core-property metadata (author, title, ...).

    Regions with no safe first-class rewrite path (comments, footnotes,
    endnotes, textboxes) FAIL CLOSED: if any carries real text we raise
    ``PartialRedactionError`` and emit NOTHING."""
    import docx

    # Fail closed BEFORE doing any work if the doc carries text we can't rewrite.
    unhandled = _docx_unhandled_text_regions(data)
    if unhandled:
        raise PartialRedactionError(
            "docx has text-bearing region(s) that cannot be safely redacted: "
            + ", ".join(sorted(set(unhandled)))
        )

    report = _new_report("docx")
    regions: set[str] = set()
    document = docx.Document(io.BytesIO(data))

    def _walk_tables(tables, prefix: str, region: str) -> None:
        for ti, table in enumerate(tables):
            for ri, row in enumerate(table.rows):
                for ci, cell in enumerate(row.cells):
                    loc = {"region": region, "table": f"{prefix}{ti}",
                           "row": ri, "cell": ci}
                    for paragraph in cell.paragraphs:
                        _redact_paragraph(paragraph, redact_fn, report, loc)
                    if cell.tables:
                        _walk_tables(cell.tables, f"{prefix}{ti}.{ri}.{ci}.", region)

    # 1) Body.
    for i, paragraph in enumerate(document.paragraphs):
        _redact_paragraph(paragraph, redact_fn, report,
                          {"region": "body", "paragraph": i})
    _walk_tables(document.tables, "", "body")
    regions.add("body")

    # 2) Headers / footers, all six variants, across every section. Linked
    #    header/footers inherit a prior section's part (no own definition); skip
    #    them, and de-dupe by part identity so a shared part is never processed
    #    twice.
    hf_attrs = (
        "header", "footer",
        "first_page_header", "first_page_footer",
        "even_page_header", "even_page_footer",
    )
    processed_parts: set[int] = set()
    for si, section in enumerate(document.sections):
        for attr in hf_attrs:
            hf = getattr(section, attr, None)
            if hf is None or hf.is_linked_to_previous:
                continue
            try:
                pid = id(hf.part)
            except Exception:
                pid = None
            if pid is not None:
                if pid in processed_parts:
                    continue
                processed_parts.add(pid)
            for i, paragraph in enumerate(hf.paragraphs):
                _redact_paragraph(paragraph, redact_fn, report,
                                  {"region": attr, "section": si, "paragraph": i})
            _walk_tables(hf.tables, f"{attr}:", attr)
            regions.add("header" if "header" in attr else "footer")

    # 3) Core-property metadata.
    _redact_metadata(document, redact_fn, report, regions)

    report["regions"] = sorted(regions)
    out = io.BytesIO()
    document.save(out)
    return DocumentRedactionResult(out.getvalue(), "docx", report)


# --- dispatch ---------------------------------------------------------------
_PDF_MAGIC = b"%PDF-"
# DOCX is a ZIP; OOXML zips start with "PK\x03\x04".
_ZIP_MAGIC = b"PK\x03\x04"


def sniff_format(data: bytes, filename: str | None, content_type: str | None) -> str:
    """Best-effort format detection: magic bytes first, then filename/mime."""
    if data[:5] == _PDF_MAGIC:
        return "pdf"
    if data[:4] == _ZIP_MAGIC:
        # A .docx is a zip; confirm via the extension/mime to avoid treating an
        # arbitrary zip as a document.
        name = (filename or "").lower()
        ct = (content_type or "").lower()
        if name.endswith(".docx") or "wordprocessingml" in ct:
            return "docx"
    name = (filename or "").lower()
    ct = (content_type or "").lower()
    if name.endswith(".pdf") or ct == "application/pdf":
        return "pdf"
    if name.endswith(".docx") or "wordprocessingml" in ct:
        return "docx"
    raise UnsupportedDocument("unsupported_document_type")


def redact_document(
    data: bytes,
    redact_fn: RedactFn,
    *,
    filename: str | None = None,
    content_type: str | None = None,
) -> DocumentRedactionResult:
    """Dispatch to the PDF or DOCX redactor based on content sniffing.

    ``redact_fn`` is the caller-bound detector (``engine.redact`` with the loaded
    analyzer/PHI models). If it raises ``RedactionUnavailable`` we do NOT catch
    it here: the exception propagates so the endpoint fails closed and emits no
    partially-redacted document.
    """
    fmt = sniff_format(data, filename, content_type)
    if fmt == "pdf":
        return redact_pdf(data, redact_fn)
    return redact_docx(data, redact_fn)


# ===========================================================================
# Human-in-the-loop: PROPOSE (detect only, no burn-in) + APPLY (burn the
# human-approved subset). The one-shot redact_document above detects AND burns
# in a single pass; these two split it so a UI can show proposed redactions,
# take human overrides (remove a false positive, add a missed span), then
# finalize. Detection logic is UNCHANGED — propose reuses the SAME word->rect
# mapping the burn-in path uses; apply reuses the SAME burn-in primitives.
# ===========================================================================


# SHARED SPAN CONTRACT (the API + web build to this exactly):
#   span = {id, type, start, end, text, confidence, source, page, rects}
# where source in {regex, presidio, deid, llm, medical, term}, confidence 0..1,
# and page + rects are the PDF page index + overlay rectangles (null for DOCX).
def _span_id(page: int | None, start: int, end: int, etype: str) -> str:
    """Stable, position-seeded id: ``"{page}:{start}:{end}:{type}"``.

    Including ``page`` disambiguates two PDF spans that share per-page char
    offsets on different pages; for DOCX ``page`` is ``None`` and the offsets are
    document-global, so the id is still unique."""
    return f"{page}:{start}:{end}:{etype}"


def _to_contract_span(
    ent: dict[str, Any],
    *,
    start: int,
    end: int,
    page: int | None,
    rects: list[list[float]] | None,
    text: str,
) -> dict[str, Any]:
    """Enrich a raw detector span into the FULL shared span contract dict.

    ``start``/``end`` are the offsets to publish (per-page for PDF, document-
    global for DOCX); ``text`` is the verbatim slice; ``page``/``rects`` are the
    overlay geometry (null for DOCX/text)."""
    etype = ent["type"]
    return {
        "id": _span_id(page, start, end, etype),
        "type": etype,
        "start": start,
        "end": end,
        "text": ent.get("text", text),
        "confidence": float(ent.get("score", 1.0)),
        "source": ent.get("source", "regex"),
        "page": page,
        "rects": rects,
    }


# --- PROPOSE: PDF -----------------------------------------------------------
def _propose_pdf(data: bytes, detect_fn: DetectFn) -> dict[str, Any]:
    """Detect PII in a PDF and return proposals with page + rects, WITHOUT
    burning anything in. Uses the exact word->rect mapping redact_pdf uses, so a
    later apply_document can burn the approved rects directly."""
    import fitz

    doc = fitz.open(stream=data, filetype="pdf")
    proposed: list[dict[str, Any]] = []
    try:
        page_count = doc.page_count
        for pno in range(page_count):
            page = doc.load_page(pno)
            words = page.get_text("words")
            if not words and ocr_available():
                words = _ocr_words(page)
            if not words:
                continue
            page_text, word_spans = _build_page_text_from_words(words)
            for ent in detect_fn(page_text):  # may raise RedactionUnavailable
                rects = _rects_for_entity(ent, word_spans)
                if not rects:
                    # No covering word rect -> cannot be shown/burned; skip it,
                    # exactly as redact_pdf skips a span with no rects.
                    continue
                rect_list = [[r.x0, r.y0, r.x1, r.y1] for r in rects]
                proposed.append(
                    _to_contract_span(
                        ent, start=ent["start"], end=ent["end"], page=pno,
                        rects=rect_list,
                        text=ent.get("text", page_text[ent["start"]:ent["end"]]),
                    )
                )
        return {"format": "pdf", "page_count": page_count,
                "proposed_spans": proposed}
    finally:
        doc.close()


# --- PROPOSE / APPLY: DOCX unit walk ---------------------------------------
_HF_ATTRS = (
    "header", "footer",
    "first_page_header", "first_page_footer",
    "even_page_header", "even_page_footer",
)


def _region_label(locator: dict[str, Any]) -> str:
    """Collapse a unit locator to the coarse region name used in report.regions
    (body / header / footer / metadata), matching redact_docx's semantics."""
    region = str(locator.get("region", ""))
    if region == "metadata":
        return "metadata"
    if "header" in region:
        return "header"
    if "footer" in region:
        return "footer"
    return "body"


def _write_paragraph(paragraph, new_text: str) -> None:
    """Overwrite a paragraph's text in the part XML (no overlay), preserving the
    first run's formatting — the same run-rewrite _redact_paragraph performs."""
    runs = paragraph.runs
    if runs:
        runs[0].text = new_text
        for r in runs[1:]:
            r.text = ""
    else:
        paragraph.add_run(new_text)


def _iter_docx_text_units(document):
    """Yield ``(locator, get_text, set_text)`` for every safely-rewritable text
    region of a DOCX, in a STABLE order shared by propose AND apply.

    Because propose and apply both walk this identical order over the identical
    source bytes, the document-global char offsets they compute line up exactly,
    so an approved span's [start, end) unambiguously maps back to one unit. Order:
    body paragraphs -> body tables (recursive) -> each section's header/footer
    variants -> core-property metadata. Mirrors redact_docx's coverage; the
    unrewritable regions (comments/footnotes/endnotes/textboxes) are handled by
    the shared fail-closed check BEFORE this runs."""
    units: list[tuple[dict[str, Any], Callable[[], str], Callable[[str], None]]] = []

    def _para_unit(paragraph, locator):
        units.append(
            (locator, (lambda p=paragraph: p.text),
             (lambda new_text, p=paragraph: _write_paragraph(p, new_text)))
        )

    def _walk_tables(tables, prefix: str, region: str) -> None:
        for ti, table in enumerate(tables):
            for ri, row in enumerate(table.rows):
                for ci, cell in enumerate(row.cells):
                    loc = {"region": region, "table": f"{prefix}{ti}",
                           "row": ri, "cell": ci}
                    for pi, paragraph in enumerate(cell.paragraphs):
                        _para_unit(paragraph, {**loc, "paragraph": pi})
                    if cell.tables:
                        _walk_tables(cell.tables, f"{prefix}{ti}.{ri}.{ci}.", region)

    # 1) Body.
    for i, paragraph in enumerate(document.paragraphs):
        _para_unit(paragraph, {"region": "body", "paragraph": i})
    _walk_tables(document.tables, "", "body")

    # 2) Headers / footers (all variants, de-duped by part identity).
    processed_parts: set[int] = set()
    for si, section in enumerate(document.sections):
        for attr in _HF_ATTRS:
            hf = getattr(section, attr, None)
            if hf is None or hf.is_linked_to_previous:
                continue
            try:
                pid = id(hf.part)
            except Exception:
                pid = None
            if pid is not None:
                if pid in processed_parts:
                    continue
                processed_parts.add(pid)
            for i, paragraph in enumerate(hf.paragraphs):
                _para_unit(paragraph, {"region": attr, "section": si,
                                       "paragraph": i})
            _walk_tables(hf.tables, f"{attr}:", attr)

    # 3) Core-property metadata (free-text fields only).
    cp = document.core_properties
    for field in _METADATA_FIELDS:
        val = getattr(cp, field, None)
        if not isinstance(val, str):
            continue
        units.append((
            {"region": "metadata", "field": field},
            (lambda f=field: getattr(cp, f) or ""),
            (lambda new_text, f=field: setattr(cp, f, new_text)),
        ))

    return units


def _propose_docx(data: bytes, detect_fn: DetectFn) -> dict[str, Any]:
    """Detect PII across every rewritable DOCX region and return proposals with
    document-global char offsets (page/rects null). Fails closed on unrewritable
    text regions, exactly like redact_docx."""
    import docx

    unhandled = _docx_unhandled_text_regions(data)
    if unhandled:
        raise PartialRedactionError(
            "docx has text-bearing region(s) that cannot be safely redacted: "
            + ", ".join(sorted(set(unhandled)))
        )

    document = docx.Document(io.BytesIO(data))
    proposed: list[dict[str, Any]] = []
    base = 0
    for _locator, get_text, _set in _iter_docx_text_units(document):
        utext = get_text()
        if utext.strip():
            for ent in detect_fn(utext):  # may raise RedactionUnavailable
                gstart, gend = base + ent["start"], base + ent["end"]
                proposed.append(
                    _to_contract_span(
                        ent, start=gstart, end=gend, page=None, rects=None,
                        text=ent.get("text", utext[ent["start"]:ent["end"]]),
                    )
                )
        base += len(utext) + 1  # +1 for the "\n" separator between units
    return {"format": "docx", "page_count": None, "proposed_spans": proposed}


def propose_document(
    data: bytes,
    redact_fn_detect: DetectFn,
    *,
    filename: str | None = None,
    content_type: str | None = None,
) -> dict[str, Any]:
    """PROPOSE: detect PII and return ``{format, page_count, proposed_spans}``
    WITHOUT altering the document. Each proposed span is the full shared contract
    dict (id/type/start/end/text/confidence/source/page/rects). ``redact_fn_detect``
    is the caller-bound detect-only function (engine.detect with loaded models);
    if it raises ``RedactionUnavailable`` (strict, NER down) we propagate so the
    endpoint fails closed and proposes NOTHING."""
    fmt = sniff_format(data, filename, content_type)
    if fmt == "pdf":
        return _propose_pdf(data, redact_fn_detect)
    return _propose_docx(data, redact_fn_detect)


# --- APPLY: PDF -------------------------------------------------------------
def _apply_pdf(data: bytes, approved_spans: list[dict[str, Any]]) -> DocumentRedactionResult:
    """Burn in EXACTLY the approved spans on a PDF. Uses each span's ``rects``
    when present (the reviewed geometry); otherwise re-maps from its per-page
    char range with the same word->rect mapping propose used. Honors human
    overrides: only the spans passed here are redacted."""
    import fitz

    report = _new_report("pdf")
    by_page: dict[Any, list[dict[str, Any]]] = defaultdict(list)
    for s in approved_spans:
        by_page[s.get("page")].append(s)

    doc = fitz.open(stream=data, filetype="pdf")
    try:
        for pno in range(doc.page_count):
            page_spans = by_page.get(pno, [])
            if not page_spans:
                continue
            page = doc.load_page(pno)
            word_spans: list[tuple[int, int, Any]] | None = None  # lazily built
            for s in page_spans:
                rects = s.get("rects")
                if not rects:
                    # No reviewed geometry (e.g. a human-ADDED span given by char
                    # range): re-map from start/end using the same mapping propose
                    # used, so an added span still burns in.
                    if word_spans is None:
                        words = page.get_text("words")
                        if not words and ocr_available():
                            words = _ocr_words(page)
                        _, word_spans = _build_page_text_from_words(words or [])
                    frects = _rects_for_entity(s, word_spans)
                    rects = [[r.x0, r.y0, r.x1, r.y1] for r in frects]
                if not rects:
                    continue
                _tally(report, s["type"])
                loc_rects: list[list[float]] = []
                for r in rects:
                    rect = fitz.Rect(r[0], r[1], r[2], r[3])
                    page.add_redact_annot(rect, fill=(0, 0, 0))
                    loc_rects.append([rect.x0, rect.y0, rect.x1, rect.y1])
                report["locations"].append(
                    {"page": pno, "type": s["type"], "rects": loc_rects}
                )
            page.apply_redactions()  # DESTROY glyphs under every annot on the page
        out = io.BytesIO()
        doc.save(out, garbage=4, deflate=True, clean=True)
        return DocumentRedactionResult(out.getvalue(), "pdf", report)
    finally:
        doc.close()


# --- APPLY: DOCX ------------------------------------------------------------
def _splice_tokens(
    text: str,
    local_spans: list[tuple[int, int, str]],
    report: dict[str, Any],
    locator: dict[str, Any],
    regions: set[str],
) -> str:
    """Overwrite only the approved [start, end) ranges of ``text`` with
    ``[REDACTED_<TYPE>_<NNN>]`` tokens (the same token format the one-shot DOCX
    path emits). ``local_spans`` are unit-local (start, end, type); overlaps are
    resolved earliest-start/longest-first so nothing is double-emitted."""
    local_spans = sorted(local_spans, key=lambda s: (s[0], -(s[1] - s[0])))
    counters: dict[str, int] = {}
    out: list[str] = []
    cursor = 0
    for st, en, ty in local_spans:
        if st < cursor:  # overlapped by an already-emitted span
            continue
        counters[ty] = counters.get(ty, 0) + 1
        out.append(text[cursor:st])
        out.append(f"[REDACTED_{ty}_{counters[ty]:03d}]")
        cursor = en
        _tally(report, ty)
        report["locations"].append({**locator, "type": ty})
        regions.add(_region_label(locator))
    out.append(text[cursor:])
    return "".join(out)


def _apply_docx(data: bytes, approved_spans: list[dict[str, Any]]) -> DocumentRedactionResult:
    """Burn in EXACTLY the approved spans on a DOCX by REPLACING the approved
    text ranges (no overlay). Maps each span's document-global [start, end) back
    to its unit via the SAME stable walk propose used. Fails closed on
    unrewritable regions, exactly like redact_docx."""
    import docx

    unhandled = _docx_unhandled_text_regions(data)
    if unhandled:
        raise PartialRedactionError(
            "docx has text-bearing region(s) that cannot be safely redacted: "
            + ", ".join(sorted(set(unhandled)))
        )

    report = _new_report("docx")
    regions: set[str] = set()
    document = docx.Document(io.BytesIO(data))

    base = 0
    for locator, get_text, set_text in _iter_docx_text_units(document):
        utext = get_text()
        ulen = len(utext)
        lo, hi = base, base + ulen
        local: list[tuple[int, int, str]] = []
        for s in approved_spans:
            gs, ge = int(s["start"]), int(s["end"])
            if ge > gs and gs >= lo and ge <= hi:
                local.append((gs - base, ge - base, s["type"]))
        if local:
            set_text(_splice_tokens(utext, local, report, locator, regions))
        base += ulen + 1  # +1 for the "\n" separator between units

    report["regions"] = sorted(regions)
    out = io.BytesIO()
    document.save(out)
    return DocumentRedactionResult(out.getvalue(), "docx", report)


def apply_document(
    data: bytes,
    approved_spans: list[dict[str, Any]],
    *,
    filename: str | None = None,
    content_type: str | None = None,
) -> dict[str, Any]:
    """APPLY: burn in EXACTLY ``approved_spans`` and return
    ``{format, redacted_document, report}``.

    This honors human overrides completely: a span the reviewer removed is simply
    absent from ``approved_spans`` and is NOT redacted; a span the reviewer added
    IS redacted (PDF via its rects or a re-map from its char range, DOCX via its
    text range). apply does NOT re-run detection, so the strict-mode fail-closed
    decision happens upstream at propose time; nothing is ever emitted from a
    detection that failed closed, because that request produces no proposals to
    approve."""
    fmt = sniff_format(data, filename, content_type)
    if fmt == "pdf":
        result = _apply_pdf(data, approved_spans)
    else:
        result = _apply_docx(data, approved_spans)
    return {"format": result.fmt, "redacted_document": result.content,
            "report": result.report}
